#!/usr/bin/env python3
"""
file_watcher_daemon.py — File-system watcher daemon for Presence

Mirrors realtime_listener.py structure:
  - Flask server on port 5556 with /pick-folder and /status endpoints
  - watchdog monitors selected folder for file changes
  - Writes events to activity_signal table (source: 'file_watcher')
  - Same env vars: SUPABASE_URL, SUPABASE_KEY (via .env / dotenv)

Usage:
  python3 file_watcher_daemon.py                # start daemon
  curl http://localhost:5556/pick-folder         # open macOS folder picker
  curl http://localhost:5556/status              # check status
"""

import subprocess
import threading
import time
import logging
import os
import sys
import json
import difflib
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
import requests
from docx import Document
import openpyxl
import pdfplumber
from flask import Flask, jsonify
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

load_dotenv()

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------
FLASK_PORT = 5556
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_KEY") or os.getenv("SUPABASE_SERVICE_ROLE_KEY") or os.getenv("SUPABASE_KEY")

BASE_DIR = Path(__file__).parent

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger("file_watcher_daemon")

# ---------------------------------------------------------------------------
# STATE
# ---------------------------------------------------------------------------
watched_folder: str | None = None
observer: Observer | None = None
running = True
events_written = 0
last_event_at: float = 0.0
started_at = time.time()

_HEARTBEAT_LAST_SENT: dict[str, float] = {}
HEARTBEAT_INTERVAL_SECONDS = 5 * 60
FILE_CONTENT_CACHE: dict[str, str] = {}


# ---------------------------------------------------------------------------
# HEARTBEAT (mirrors realtime_listener.py)
# ---------------------------------------------------------------------------
def beat_heartbeat(agent_name: str) -> None:
    """Best-effort heartbeat to Supabase RPC. Non-fatal and self-throttled."""
    now = time.time()
    last_sent = _HEARTBEAT_LAST_SENT.get(agent_name, 0.0)
    if (now - last_sent) < HEARTBEAT_INTERVAL_SECONDS:
        return

    supabase_url = (SUPABASE_URL or "").strip().rstrip("/")
    supabase_key = (SUPABASE_KEY or "").strip()
    if not supabase_url or not supabase_key:
        return

    url = f"{supabase_url}/rest/v1/rpc/beat"
    headers = {
        "apikey": supabase_key,
        "Authorization": f"Bearer {supabase_key}",
        "Content-Type": "application/json",
    }
    try:
        requests.post(
            url,
            json={"p_name": agent_name, "p_meta": {}},
            headers=headers,
            timeout=10,
        )
        _HEARTBEAT_LAST_SENT[agent_name] = now
    except Exception as exc:
        logger.warning(f"[Heartbeat] {agent_name} beat failed: {exc}")


# ---------------------------------------------------------------------------
# ACTIVITY SIGNAL (mirrors realtime_listener.py write_activity_signal)
# ---------------------------------------------------------------------------
def write_activity_signal(
    signal_type: str,
    signal_value: str,
    metadata: dict | None = None,
    user_id: str = "default",
) -> bool:
    """Best-effort write to activity_signal. Non-fatal."""
    try:
        payload = {
            "signal_type": signal_type,
            "signal_value": signal_value,
            "metadata": metadata or {},
            "user_id": user_id,
        }
        resp = requests.post(
            f"{SUPABASE_URL}/rest/v1/activity_signal",
            json=payload,
            headers={
                "apikey": SUPABASE_KEY,
                "Authorization": f"Bearer {SUPABASE_KEY}",
                "Content-Type": "application/json",
                "Prefer": "return=minimal",
            },
            timeout=10,
        )
        if resp.status_code not in (200, 201, 204):
            print(
                f"[FW:Signal] Write failed: {resp.status_code} {resp.text[:200]}",
                file=sys.stderr,
            )
            return False
        return True
    except Exception as exc:
        print(f"[FW:Signal] Error writing activity_signal: {exc}", file=sys.stderr)
        return False


# ---------------------------------------------------------------------------
# WATCHDOG HANDLER
# ---------------------------------------------------------------------------
class PresenceFileHandler(FileSystemEventHandler):
    """Writes file-change events to activity_signal with source: 'file_watcher'."""

    # Debounce: skip duplicate events for the same path within this window
    DEBOUNCE_SECONDS = 2

    def __init__(self):
        super().__init__()
        self._recent: dict[str, float] = {}

    def _debounce(self, path: str) -> bool:
        now = time.time()
        last = self._recent.get(path, 0.0)
        if now - last < self.DEBOUNCE_SECONDS:
            return True
        self._recent[path] = now
        # Prune old entries
        cutoff = now - 30
        self._recent = {k: v for k, v in self._recent.items() if v > cutoff}
        return False

    def _extract_text(self, path: str) -> str | None:
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext in {".txt", ".md"}:
                with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                    return fh.read()
            elif ext == ".docx":
                doc = Document(path)
                return "\n".join(paragraph.text for paragraph in doc.paragraphs)
            elif ext == ".xlsx":
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    lines = []
                    for ws in wb.worksheets:
                        for row in ws.iter_rows(values_only=True):
                            values = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                            if values:
                                lines.append(" | ".join(values))
                    return "\n".join(lines)
                finally:
                    wb.close()
            elif ext == ".pdf":
                with pdfplumber.open(path) as pdf:
                    return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except FileNotFoundError:
            return ""
        except IsADirectoryError:
            return ""
        except Exception:
            return None
        return None

    def _emit_signal(self, event_type: str, src: str, rel_path: str, metadata: dict) -> None:
        if event_type == "deleted":
            FILE_CONTENT_CACHE.pop(src, None)
            write_activity_signal(
                signal_type="file_change",
                signal_value=f"{event_type}: {rel_path}",
                metadata=metadata,
                user_id="default",
            )
            return

        new_text = self._extract_text(src)
        if new_text is None:
            snippet = "[binary file changed]"
        else:
            diff = difflib.ndiff(FILE_CONTENT_CACHE.get(src, "").splitlines(), new_text.splitlines())
            diff_lines = [line[2:] for line in diff if line.startswith("+ ")]
            if not diff_lines:
                return
            FILE_CONTENT_CACHE[src] = new_text
            snippet = "\n".join(diff_lines).strip()
            if not snippet:
                return

        metadata["snippet"] = snippet[:500]
        wrote = write_activity_signal(
            signal_type="file_change",
            signal_value=f"{event_type}: {rel_path} | {metadata['snippet']}",
            metadata=metadata,
            user_id="default",
        )
        if wrote:
            global events_written, last_event_at
            events_written += 1
            last_event_at = time.time()
            logger.info(f"[FW:Event] {event_type}: {rel_path}")

    def on_any_event(self, event):
        # Skip directory-level events and temporary/hidden files
        if event.is_directory:
            return
        src = getattr(event, "src_path", "")
        if not src:
            return
        basename = os.path.basename(src)
        if basename.startswith(".") or basename.startswith("~"):
            return

        if self._debounce(src):
            return

        event_type = event.event_type  # created, modified, deleted, moved
        rel_path = src
        if watched_folder:
            try:
                rel_path = os.path.relpath(src, watched_folder)
            except ValueError:
                pass

        ext = os.path.splitext(src)[1].lower()
        signal_value = f"{event_type}: {rel_path}"

        metadata = {
            "source": "file_watcher",
            "event_type": event_type,
            "file_path": rel_path,
            "extension": ext,
            "watched_folder": watched_folder or "",
        }

        # For move events, include the destination
        if event_type == "moved":
            dest = getattr(event, "dest_path", "")
            if dest and watched_folder:
                try:
                    dest = os.path.relpath(dest, watched_folder)
                except ValueError:
                    pass
            metadata["dest_path"] = dest
            signal_value = f"{event_type}: {rel_path} -> {dest}"

        self._emit_signal(event_type, src, rel_path if event_type != "moved" else signal_value, metadata)


# ---------------------------------------------------------------------------
# FOLDER PICKER (macOS osascript)
# ---------------------------------------------------------------------------
def pick_folder_dialog() -> str | None:
    """Open native macOS folder picker via osascript. Returns path or None."""
    script = (
        'tell application "System Events"\n'
        '  activate\n'
        '  set chosenFolder to choose folder with prompt "Select a folder to watch"\n'
        '  return POSIX path of chosenFolder\n'
        'end tell'
    )
    try:
        result = subprocess.run(
            ["osascript", "-e", script],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()
        return None
    except Exception as exc:
        logger.error(f"[FW:Picker] osascript error: {exc}")
        return None


def start_watching(folder: str) -> bool:
    """Start watchdog observer on the given folder."""
    global observer, watched_folder

    # Stop existing observer if any
    if observer is not None:
        try:
            observer.stop()
            observer.join(timeout=5)
        except Exception:
            pass

    if not os.path.isdir(folder):
        logger.error(f"[FW:Watch] Not a directory: {folder}")
        return False

    watched_folder = folder
    handler = PresenceFileHandler()
    observer = Observer()
    observer.schedule(handler, folder, recursive=True)
    observer.start()
    logger.info(f"[FW:Watch] Now watching: {folder}")

    # Write a signal that watching started
    write_activity_signal(
        signal_type="file_watcher_started",
        signal_value=f"Watching: {folder}",
        metadata={"source": "file_watcher", "watched_folder": folder},
        user_id="default",
    )
    return True


def stop_watching() -> bool:
    """Stop watchdog observer and clear watched folder state."""
    global observer, watched_folder
    try:
        if observer is not None:
            observer.stop()
            observer.join(timeout=5)
        observer = None
        previous = watched_folder
        watched_folder = None
        write_activity_signal(
            signal_type="file_watcher_stopped",
            signal_value="Watching stopped",
            metadata={"source": "file_watcher", "watched_folder": previous or ""},
            user_id="default",
        )
        return True
    except Exception as exc:
        logger.error(f"[FW:Stop] Failed to stop watcher: {exc}")
        return False


# ---------------------------------------------------------------------------
# FLASK SERVER
# ---------------------------------------------------------------------------
app = Flask(__name__)


@app.route("/status")
def status():
    return jsonify({
        "status": "running",
        "watched_folder": watched_folder,
        "observer_alive": observer.is_alive() if observer else False,
        "events_written": events_written,
        "last_event_at": datetime.fromtimestamp(last_event_at, tz=timezone.utc).isoformat() if last_event_at else None,
        "uptime_seconds": int(time.time() - started_at),
    })


@app.route("/pick-folder", methods=["GET", "POST"])
def pick_folder():
    folder = pick_folder_dialog()
    if folder is None:
        return jsonify({"error": "No folder selected or dialog cancelled"}), 400

    ok = start_watching(folder)
    if not ok:
        return jsonify({"error": f"Failed to watch: {folder}"}), 500

    return jsonify({"watched_folder": folder, "status": "watching"})


@app.route("/stop-watching", methods=["POST"])
def stop_watching_route():
    ok = stop_watching()
    if not ok:
        return jsonify({"ok": False, "error": "Failed to stop watcher"}), 500
    return jsonify({"ok": True, "watched_folder": None, "status": "stopped"})


# ---------------------------------------------------------------------------
# PLIST MANAGEMENT
# ---------------------------------------------------------------------------
PLIST_LABEL = "com.presence.filewatcher"
PLIST_PATH = Path.home() / "Library" / "LaunchAgents" / f"{PLIST_LABEL}.plist"


def create_and_load_plist() -> None:
    """Create launchd plist mirroring com.presence.realtime-listener and load it."""
    script_path = os.path.abspath(__file__)
    working_dir = os.path.dirname(script_path)
    python_path = "/opt/homebrew/bin/python3"
    log_dir = Path.home() / "Library" / "Logs"

    plist_content = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE plist PUBLIC "-//Apple//DTD PLIST 1.0//EN" "http://www.apple.com/DTDs/PropertyList-1.0.dtd">
<plist version="1.0">
<dict>
  <key>Label</key>
  <string>{PLIST_LABEL}</string>

  <key>ProgramArguments</key>
  <array>
    <string>{python_path}</string>
    <string>{script_path}</string>
  </array>

  <key>WorkingDirectory</key>
  <string>{working_dir}</string>

  <key>RunAtLoad</key>
  <true/>

  <key>KeepAlive</key>
  <true/>

  <key>StandardOutPath</key>
  <string>{log_dir}/presence-filewatcher.log</string>

  <key>StandardErrorPath</key>
  <string>{log_dir}/presence-filewatcher-error.log</string>

  <key>EnvironmentVariables</key>
  <dict>
    <key>PATH</key>
    <string>/opt/homebrew/bin:/usr/local/bin:/usr/bin:/bin:/usr/sbin:/sbin</string>
  </dict>
</dict>
</plist>
"""

    # Unload existing if present
    if PLIST_PATH.exists():
        try:
            subprocess.run(
                ["launchctl", "unload", str(PLIST_PATH)],
                capture_output=True,
                timeout=10,
            )
        except Exception:
            pass

    PLIST_PATH.parent.mkdir(parents=True, exist_ok=True)
    PLIST_PATH.write_text(plist_content)
    logger.info(f"[FW:Plist] Written: {PLIST_PATH}")

    # Load plist
    result = subprocess.run(
        ["launchctl", "load", str(PLIST_PATH)],
        capture_output=True,
        text=True,
        timeout=10,
    )
    if result.returncode == 0:
        logger.info(f"[FW:Plist] Loaded: {PLIST_LABEL}")
    else:
        logger.error(f"[FW:Plist] Load failed: {result.stderr.strip()}")


# ---------------------------------------------------------------------------
# ENTRY POINT
# ---------------------------------------------------------------------------
def main():
    global running

    # Validate env
    if not SUPABASE_URL or not SUPABASE_KEY:
        print("ERROR: SUPABASE_URL and SUPABASE_SERVICE_KEY (or SUPABASE_KEY) must be set in .env or environment", file=sys.stderr)
        sys.exit(1)

    # --install flag: write plist and exit (the plist will start this script)
    if "--install" in sys.argv:
        create_and_load_plist()
        # Verify the daemon came up
        time.sleep(2)
        try:
            r = requests.get(f"http://localhost:{FLASK_PORT}/status", timeout=5)
            if r.status_code == 200:
                logger.info(f"[FW:Verify] Daemon running: {r.json()}")
            else:
                logger.warning(f"[FW:Verify] /status returned {r.status_code}")
        except Exception as exc:
            logger.warning(f"[FW:Verify] Could not reach daemon: {exc}")
        return

    print(f"""
╔══════════════════════════════════════════════╗
║  Presence File Watcher Daemon                ║
║  Port: {FLASK_PORT}                                  ║
║  Supabase: {(SUPABASE_URL or '')[:30]:30s}   ║
║  Ctrl+C to stop                              ║
╚══════════════════════════════════════════════╝
""")

    # Start heartbeat thread
    def heartbeat_loop():
        while running:
            beat_heartbeat("file_watcher_daemon")
            time.sleep(1)

    hb_thread = threading.Thread(target=heartbeat_loop, daemon=True, name="heartbeat")
    hb_thread.start()

    # Start Flask (blocks)
    try:
        app.run(host="127.0.0.1", port=FLASK_PORT, debug=False)
    except KeyboardInterrupt:
        print("\n[FW] Shutting down...")
        running = False
        if observer:
            observer.stop()
            observer.join(timeout=5)
        print("[FW] Stopped.")


if __name__ == "__main__":
    main()
